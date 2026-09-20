# -*- coding: utf-8 -*-
# part of api/devices/devices_routes.py (D2 split)

import logging
import time
import json
import threading
import io
import openpyxl
from flask_restx import Namespace, Resource, fields
from models import Device, DeviceHeartbeat, ClassInfo, Admin, get_by_id
from sqlalchemy.orm import joinedload
from utils.permission import requires_permission, get_current_admin, get_admin_class_ids
from utils.response import APIResponse
from utils.decorators import safe_handle
from utils.pagination import get_pagination
from services.heartbeat_service import is_device_online
from services.device_service import (
    create_device,
    update_device,
    delete_device,
    bind_device_class,
    bind_device_admin,
    resolve_device_alert,
    update_device_settings,
    import_devices,
    revoke_device_secret,
)
from services.device_query_service import (
    get_device_list_view,
    get_device_alerts_view,
    get_device_stats_view,
    get_device_advanced_stats_view,
)
from utils.api_cache_middleware import cached_api, invalidate_cache
from datetime import datetime, timedelta
from models import db
from sqlalchemy import func
from models import Alert
from flask import request
from flask import send_file
from flask import Response

import api.devices.devices_routes as _mod

from api.devices.devices_routes import logger, ns_devices, send_ota_upgrade_command, device_model, device_list_response, device_stats_response, bind_class_model, bind_admin_model, _UNLOCK_SPECS, _restart_command_payload, _send_device_restart, _publish_unlock_retry, _start_smart_unlock, ota_upgrade_model

@ns_devices.route("/<int:id>/ota-upgrade")
@ns_devices.param("id", "设备ID")
class DeviceOTAUpgrade(Resource):

    @ns_devices.doc("device_ota_upgrade", description="设备OTA固件升级", security="Bearer")
    @ns_devices.expect(ota_upgrade_model)
    @ns_devices.response(200, "成功")
    @ns_devices.response(400, "设备不在线")
    @requires_permission("device.edit")
    def post(self, id):
        """
        设备OTA固件升级

        向指定设备发送OTA固件升级指令。
        需要设备在线才能执行升级。

        请求体：
        - firmware_url: 固件下载URL（必填）
        - version: 目标固件版本（可选）
        - force: 是否强制升级，忽略版本检查（可选，默认false）
        """
        device = Device.query.get_or_404(id)

        # 差异 #10：统一走 is_device_online（last_heartbeat 时效性）。
        # 向后兼容：保留 status == "online" 兜底，避免无心跳的既有数据被判离线。
        if not (is_device_online(device) or device.status == "online"):
            return APIResponse.bad_request(message="设备不在线，无法执行OTA升级")

        data = request.get_json()
        firmware_url = data.get("firmware_url")
        version = data.get("version", "")
        force = data.get("force", False)

        return send_ota_upgrade_command(firmware_url, version, force, device.device_id)

@ns_devices.route("/ota-upgrade-all")
class DeviceOTAUpgradeAll(Resource):

    @ns_devices.doc("device_ota_upgrade_all", description="批量OTA固件升级", security="Bearer")
    @ns_devices.expect(ota_upgrade_model)
    @ns_devices.response(200, "成功")
    @requires_permission("device.edit")
    def post(self):
        """
        批量OTA固件升级

        向所有在线设备发送OTA固件升级指令。

        请求体：
        - firmware_url: 固件下载URL（必填）
        - version: 目标固件版本（可选）
        - force: 是否强制升级，忽略版本检查（可选，默认false）
        """
        data = request.get_json()
        firmware_url = data.get("firmware_url")
        version = data.get("version", "")
        force = data.get("force", False)

        return send_ota_upgrade_command(firmware_url, version, force)

@ns_devices.route("/bulk-ota-upgrade")
class DeviceBulkOTAUpgrade(Resource):

    @ns_devices.doc(
        "device_bulk_ota_upgrade", description="批量OTA固件升级（别名）", security="Bearer"
    )
    @ns_devices.expect(ota_upgrade_model)
    @ns_devices.response(200, "成功")
    @requires_permission("device.edit")
    def post(self):
        """
        批量OTA固件升级（别名接口）

        向所有在线设备发送OTA固件升级指令。
        """
        data = request.get_json()
        firmware_url = data.get("firmware_url")
        version = data.get("version", "")
        force = data.get("force", False)

        return send_ota_upgrade_command(firmware_url, version, force)

@ns_devices.route("/import")
class DeviceImport(Resource):

    @ns_devices.doc("device_import", description="批量导入设备", security="Bearer")
    @requires_permission("device.edit")
    @safe_handle(message="设备操作失败，请稍后重试", default_status=500, error_code="INTERNAL_ERROR")
    def post(self):
        """
        批量导入设备

        通过Excel文件批量导入设备信息。

        支持的字段：设备标识(device_id)、设备名称(name)、班级名称(class_name)、管理员姓名(admin_name)

        返回导入结果统计。
        """
        if "file" not in request.files:
            return APIResponse.bad_request(message="没有上传文件")

        file = request.files["file"]

        if not file.filename.endswith((".xlsx", ".xls")):
            return APIResponse.bad_request(message="仅支持Excel文件格式")

        result = import_devices(file)
        return APIResponse.success(data=result)

@ns_devices.route("/export")
class DeviceExport(Resource):

    @ns_devices.doc("device_export", description="导出设备数据", security="Bearer")
    @requires_permission("device.view")
    @safe_handle(message="设备操作失败，请稍后重试", default_status=500, error_code="INTERNAL_ERROR")
    def get(self):
        """
        导出设备数据

        支持JSON和Excel格式导出。

        参数：
        - format: 导出格式（json 或 excel，默认excel）

        返回设备数据文件下载。
        """
        format_type = request.args.get("format", "excel")

        devices = Device.query.all()

        if format_type == "json":
            device_list = []
            for device in devices:
                device_data = {
                    "id": device.id,
                    "device_id": device.device_id,
                    "name": device.name,
                    "status": device.status,
                    "is_online": is_device_online(device),
                    "last_heartbeat": (
                        device.last_heartbeat.strftime("%Y-%m-%d %H:%M:%S")
                        if device.last_heartbeat
                        else ""
                    ),
                    "wifi_signal": device.wifi_signal,
                    "uptime": device.uptime,
                    "box_a_status": device.box_a_status,
                    "box_b_status": device.box_b_status,
                    "system_state": device.system_state,
                    "class_info_id": device.class_info_id,
                    "class_name": device.class_info.name if device.class_info else "",
                    "admin_id": device.admin_id,
                    "admin_name": device.admin.name if device.admin else "",
                    "created_at": (
                        device.created_at.strftime("%Y-%m-%d %H:%M:%S")
                        if device.created_at
                        else ""
                    ),
                    "updated_at": (
                        device.updated_at.strftime("%Y-%m-%d %H:%M:%S")
                        if device.updated_at
                        else ""
                    ),
                }
                device_list.append(device_data)

            json_str = json.dumps(device_list, ensure_ascii=False, indent=2)
            return Response(
                json_str,
                mimetype="application/json",
                headers={
                    "Content-Disposition": f'attachment; filename=devices_{datetime.now().strftime("%Y%m%d_%H%M%S")}.json'
                },
            )

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "设备数据"

        headers = [
            "设备标识",
            "设备名称",
            "状态",
            "是否在线",
            "最后心跳",
            "WiFi信号",
            "班级名称",
            "管理员姓名",
            "创建时间",
            "更新时间",
        ]
        sheet.append(headers)

        header_font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        header_fill = openpyxl.styles.PatternFill(
            start_color="4A5568", end_color="4A5568", fill_type="solid"
        )
        for col in range(1, len(headers) + 1):
            cell = sheet.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill

        for device in devices:
            row_data = [
                device.device_id,
                device.name,
                device.status,
                "是" if is_device_online(device) else "否",
                (
                    device.last_heartbeat.strftime("%Y-%m-%d %H:%M:%S")
                    if device.last_heartbeat
                    else ""
                ),
                device.wifi_signal,
                device.class_info.name if device.class_info else "",
                device.admin.name if device.admin else "",
                (
                    device.created_at.strftime("%Y-%m-%d %H:%M:%S")
                    if device.created_at
                    else ""
                ),
                (
                    device.updated_at.strftime("%Y-%m-%d %H:%M:%S")
                    if device.updated_at
                    else ""
                ),
            ]
            sheet.append(row_data)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            download_name=f'devices_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx',
        )

@ns_devices.route("/<int:id>/secret")
@ns_devices.param("id", "设备ID")
class DeviceSecret(Resource):
    """设备密钥的签发 / 查看状态 / 吊销（差异 #4）。

    密钥明文**仅在签发/重置的那一次响应中返回**，此后后端只保留原值供
    HMAC 校验使用，接口不再回显 —— 避免密钥在日志/前端缓存中长期驻留。
    """

    @ns_devices.doc("rotate_device_secret", description="签发/重置设备密钥", security="Bearer")
    @ns_devices.response(200, "签发成功")
    @ns_devices.response(404, "设备不存在")
    @requires_permission("device.edit")
    def post(self, id):
        """
        签发或重置设备密钥（差异 #4 阶段 3）。

        返回的 `device_secret` 为**明文，仅此一次可见**，请立即导出烧录到设备 NVS。
        签发后该设备的上行必须携带 ts/nonce/sig（差异 #4 阶段 2 验签）。
        """
        device = get_by_id(Device, id)
        if not device:
            return APIResponse.not_found(message="设备不存在")

        from utils.device_auth import issue_device_secret

        secret = issue_device_secret(device, commit=True)
        logger.info(f"[设备认证] 已为设备 {device.device_id} 签发密钥")

        return APIResponse.success(
            message="密钥已签发，请立即导出烧录（明文仅此一次可见）",
            data={
                "id": device.id,
                "device_id": device.device_id,
                "device_secret": secret,
                "secret_issued_at": (
                    device.secret_issued_at.isoformat() if device.secret_issued_at else None
                ),
            },
        )

    @ns_devices.doc("get_device_secret_status", description="查看设备密钥状态", security="Bearer")
    @ns_devices.response(200, "成功")
    @ns_devices.response(404, "设备不存在")
    @requires_permission("device.view")
    def get(self, id):
        """查看设备密钥状态（**不返回密钥明文**）。"""
        device = get_by_id(Device, id)
        if not device:
            return APIResponse.not_found(message="设备不存在")

        return APIResponse.success(
            data={
                "id": device.id,
                "device_id": device.device_id,
                "has_secret": bool(getattr(device, "device_secret", None)),
                "secret_issued_at": (
                    device.secret_issued_at.isoformat() if device.secret_issued_at else None
                ),
                "last_seen_ts": getattr(device, "last_seen_ts", None),
            }
        )

    @ns_devices.doc("revoke_device_secret", description="吊销设备密钥", security="Bearer")
    @ns_devices.response(200, "已吊销")
    @ns_devices.response(404, "设备不存在")
    @requires_permission("device.edit")
    def delete(self, id):
        """吊销设备密钥：置空后该设备回到「未发放密钥」状态（验签放行）。

        注意：吊销**不等于封禁** —— 设备仍可上报（差异 #4 阶段 2 宽容策略）。
        若要阻止未登记设备接入，请开启系统配置 `device_whitelist_enabled`。
        """
        device = get_by_id(Device, id)
        if not device:
            return APIResponse.not_found(message="设备不存在")

        revoke_device_secret(device)
        logger.info(f"[设备认证] 已吊销设备 {device.device_id} 的密钥")

        return APIResponse.success(
            message="密钥已吊销（该设备恢复为免验签状态）",
            data={"id": device.id, "device_id": device.device_id, "has_secret": False},
        )
