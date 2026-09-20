# -*- coding: utf-8 -*-
# part of api/academics/course_schedule_routes.py (D2 split)

import json
import io
from flask_restx import Namespace, Resource, fields
from flask import request, send_file
from models import CourseSchedule, ClassInfo, Subject, ClassPeriod, Admin, ImportConfig, get_by_id
from services.academics_service import academics_service
from services.course_schedule_service import (
    get_schedule_list_view,
    get_schedule_by_class_view,
    get_schedule_now_view,
    get_schedule_options_view,
    check_schedule_conflict_view,
    build_schedule_export_data,
    _schedule_dict,
    format_day_of_week,
    get_period_info,
    check_conflicts,
    check_teacher_conflicts,
    check_classroom_conflicts,
)
from utils.permission import requires_permission, get_allowed_classes, get_current_admin
from utils.response import APIResponse
from utils.api_cache_middleware import cached_api, invalidate_cache
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import logging

from api.academics.course_schedule_routes import logger, ns_course_schedule, course_schedule_model, course_schedule_response, _DAY_TEXT_MAP, _validate_text_field, _validate_day_of_week, _validate_period_number, _validate_teacher_role, _check_schedule_update_forbidden, _resolve_schedule_update_fields, _resolve_teacher_name, _schedule_time_changed, _collect_time_conflicts, _check_change_conflicts, _resolve_final_color, _schedule_update_response

@ns_course_schedule.route("/")
class CourseScheduleList(Resource):

    @ns_course_schedule.doc("list_course_schedule", description="获取课程表列表")
    @ns_course_schedule.response(200, "成功")
    @requires_permission("schedule.view")
    @cached_api(ttl=30)
    def get(self):
        """
        获取课程表列表。非管理员用户只能查看关联班级的课程表。

        可选参数：
        - class_info_id: 班级ID
        - day_of_week: 星期
        - teacher_name: 教师姓名
        - classroom: 教室
        """
        args = ns_course_schedule.parser.parse_args()
        view = get_schedule_list_view(args)
        return APIResponse.success(
            data={"schedules": view["schedules"]},
            pagination={
                "page": view["page"],
                "per_page": view["per_page"],
                "total": view["total"],
                "pages": view["pages"],
            },
        )

    @ns_course_schedule.doc("create_course_schedule", description="创建课程安排", security="Bearer")
    @ns_course_schedule.expect(course_schedule_model)
    @ns_course_schedule.response(201, "创建成功", course_schedule_response)
    @requires_permission("schedule.manage")
    def post(self):
        """
        创建课程安排。需要课表管理权限。

        请求体：
        - class_info_id: 班级ID（必填）
        - subject_id: 科目ID（必填）
        - day_of_week: 星期(0-6)（必填）
        - period_number: 节次编号（必填）
        - teacher_name: 教师姓名（可选）
        - classroom: 教室（可选）
        - description: 描述（可选）
        - color: 颜色（可选）
        """
        data = ns_course_schedule.payload

        # 数据隔离：非管理员只能为关联班级创建课程
        admin = get_current_admin()
        allowed_classes = get_allowed_classes(admin.id) if admin else None
        if allowed_classes is not None:
            class_info = get_by_id(ClassInfo, data["class_info_id"])
            if class_info and class_info.name not in allowed_classes:
                return APIResponse.forbidden(message="无权为该班级创建课程安排")

        # 检查时间冲突
        conflicts = []
        conflicts.extend(
            check_conflicts(data["class_info_id"], data["day_of_week"], data["period_number"])
        )

        teacher_id = data.get("teacher_id")
        teacher_name = data.get("teacher_name")

        # 如果提供了teacher_id，验证教师存在并获取教师姓名
        if teacher_id:
            teacher = get_by_id(Admin, teacher_id)
            if teacher:
                teacher_name = teacher.real_name or teacher.username
            else:
                return APIResponse.bad_request(message=f'教师ID "{teacher_id}" 在系统中不存在')

        if teacher_name:
            conflicts.extend(
                check_teacher_conflicts(teacher_name, data["day_of_week"], data["period_number"])
            )

        classroom = data.get("classroom")
        if classroom:
            conflicts.extend(
                check_classroom_conflicts(classroom, data["day_of_week"], data["period_number"])
            )

        if conflicts:
            return APIResponse.bad_request(message="存在时间冲突", errors=conflicts)

        subject = get_by_id(Subject, data["subject_id"])
        color = data.get("color") or (subject.color if subject else "#3B82F6")

        schedule_id = academics_service.create_course_schedule(
            {
                "class_info_id": data["class_info_id"],
                "subject_id": data["subject_id"],
                "day_of_week": data["day_of_week"],
                "period_number": data["period_number"],
                "teacher_id": teacher_id,
                "teacher_name": teacher_name,
                "classroom": classroom,
                "description": data.get("description"),
                "color": color,
                "is_active": data.get("is_active", True),
            }
        )
        schedule = get_by_id(CourseSchedule, schedule_id)

        period_info = get_period_info(schedule.period_number)
        invalidate_cache("api:/api/course-schedules/*")
        return (
            APIResponse.success(
                data={
                    "id": schedule.id,
                    "class_info_id": schedule.class_info_id,
                    "class_name": schedule.class_info.name if schedule.class_info else "",
                    "subject_id": schedule.subject_id,
                    "subject_name": schedule.subject.name if schedule.subject else "",
                    "subject_color": schedule.subject.color if schedule.subject else schedule.color,
                    "day_of_week": schedule.day_of_week,
                    "day_of_week_text": format_day_of_week(schedule.day_of_week),
                    "period_number": schedule.period_number,
                    "period_name": period_info["name"],
                    "period_time": period_info["time"],
                    "teacher_id": schedule.teacher_id,
                    "teacher_name": schedule.teacher_name,
                    "classroom": schedule.classroom,
                    "description": schedule.description,
                    "color": schedule.color,
                    "is_active": schedule.is_active,
                    "created_at": schedule.created_at.isoformat() if schedule.created_at else None,
                },
                message="课程安排创建成功",
            ),
            201,
        )

@ns_course_schedule.route("/<int:id>")
@ns_course_schedule.param("id", "课程ID")
class CourseScheduleResource(Resource):

    @ns_course_schedule.doc("get_course_schedule", description="获取课程详情")
    @ns_course_schedule.response(200, "成功", course_schedule_response)
    @ns_course_schedule.response(404, "课程不存在")
    @requires_permission("schedule.view")
    def get(self, id):
        """获取课程详情。需要课表查看权限。"""
        schedule = CourseSchedule.query.get_or_404(id)

        # 数据隔离：非管理员只能查看关联班级的课程
        admin = get_current_admin()
        allowed_classes = get_allowed_classes(admin.id) if admin else None
        if (
            allowed_classes is not None
            and schedule.class_info
            and schedule.class_info.name not in allowed_classes
        ):
            return APIResponse.forbidden(message="无权查看该课程")

        return APIResponse.success(data=_schedule_dict(schedule))

    @ns_course_schedule.doc("update_course_schedule", description="更新课程安排", security="Bearer")
    @ns_course_schedule.expect(course_schedule_model)
    @ns_course_schedule.response(200, "更新成功")
    @ns_course_schedule.response(404, "课程不存在")
    @requires_permission("schedule.manage")
    def put(self, id):
        """更新课程安排。需要课表管理权限。"""
        schedule = CourseSchedule.query.get_or_404(id)

        # 数据隔离：非管理员只能修改关联班级的课程
        forbidden = _check_schedule_update_forbidden(schedule)
        if forbidden:
            return forbidden

        data = ns_course_schedule.payload

        # 获取更新后的字段值
        new_fields = _resolve_schedule_update_fields(data, schedule)
        new_class_info_id = new_fields["class_info_id"]
        new_day_of_week = new_fields["day_of_week"]
        new_period_number = new_fields["period_number"]
        new_teacher_id = new_fields["teacher_id"]
        new_teacher_name = new_fields["teacher_name"]
        new_classroom = new_fields["classroom"]

        # 如果提供了teacher_id，验证教师存在并获取教师姓名
        teacher_error, new_teacher_name = _resolve_teacher_name(new_teacher_id, new_teacher_name)
        if teacher_error:
            return teacher_error

        # 如果时间或班级发生变化，检查冲突
        if _schedule_time_changed(new_fields, schedule):
            conflicts = _collect_time_conflicts(
                new_class_info_id,
                new_day_of_week,
                new_period_number,
                new_teacher_name,
                new_classroom,
                id,
            )
            if conflicts:
                return APIResponse.bad_request(message="存在时间冲突", errors=conflicts)

        # 如果教师或教室发生变化，检查冲突
        change_conflict = _check_change_conflicts(
            schedule,
            new_teacher_name,
            new_classroom,
            new_day_of_week,
            new_period_number,
            new_class_info_id,
            id,
        )
        if change_conflict:
            return change_conflict

        new_subject_id = data.get("subject_id", schedule.subject_id)
        final_color = _resolve_final_color(data, new_subject_id)

        academics_service.update_course_schedule(
            id,
            {
                "class_info_id": new_class_info_id,
                "subject_id": new_subject_id,
                "day_of_week": new_day_of_week,
                "period_number": new_period_number,
                "teacher_id": new_teacher_id,
                "teacher_name": new_teacher_name,
                "classroom": new_classroom,
                "description": data.get("description", schedule.description),
                "color": final_color,
                "is_active": data.get("is_active", schedule.is_active),
            },
        )
        schedule = get_by_id(CourseSchedule, id)

        period_info = get_period_info(schedule.period_number)
        invalidate_cache("api:/api/course-schedules/*")
        return _schedule_update_response(schedule, period_info)

    @ns_course_schedule.doc("delete_course_schedule", description="删除课程安排", security="Bearer")
    @ns_course_schedule.response(200, "删除成功")
    @ns_course_schedule.response(404, "课程不存在")
    @requires_permission("schedule.manage")
    def delete(self, id):
        """删除课程安排。需要课表管理权限。"""
        schedule = CourseSchedule.query.get_or_404(id)

        # 数据隔离：非管理员只能删除关联班级的课程
        admin = get_current_admin()
        allowed_classes = get_allowed_classes(admin.id) if admin else None
        if (
            allowed_classes is not None
            and schedule.class_info
            and schedule.class_info.name not in allowed_classes
        ):
            return APIResponse.forbidden(message="无权删除该课程")

        academics_service.delete_course_schedule(id)
        invalidate_cache("api:/api/course-schedules/*")
        return APIResponse.success(message="课程安排删除成功")

@ns_course_schedule.route("/class/<int:class_info_id>")
@ns_course_schedule.param("class_info_id", "班级ID")
class CourseScheduleByClass(Resource):

    @ns_course_schedule.doc("get_course_schedule_by_class", description="获取班级课程表")
    @ns_course_schedule.response(200, "成功")
    @requires_permission("schedule.view")
    @cached_api(ttl=30)
    def get(self, class_info_id):
        """获取指定班级的完整课程表。非管理员用户只能查看关联班级的课程表。"""
        result = get_schedule_by_class_view(class_info_id)
        if result is None:
            return APIResponse.forbidden(message="无权查看该班级课程表")
        return result

@ns_course_schedule.route("/now")
class CourseScheduleNow(Resource):
    @ns_course_schedule.doc("get_course_schedule_now", description="获取当前时刻班级上课状态")
    @ns_course_schedule.response(200, "成功")
    @requires_permission("schedule.view")
    def get(self):
        """返回当前是否处于上课时间、当前节次与科目。

        可传 class_info_id 精确查询某班；也可传 device_id 按设备反查班级
        （与下发端点 _resolve_class_from_device 口径一致，供"指定设备"模式的徽章使用）。
        """
        class_info_id = request.args.get("class_info_id", type=int)
        if not class_info_id:
            device_id = request.args.get("device_id", type=str)
            if device_id:
                try:
                    from models import Device

                    dev = Device.query.filter_by(device_id=str(device_id)).first()
                    if dev and dev.class_info_id:
                        class_info_id = dev.class_info_id
                except Exception:
                    logger.warning("按设备解析 class_info_id 失败，置 None", exc_info=True)
                    class_info_id = None
        data = get_schedule_now_view(class_info_id)
        return APIResponse.success(data=data)

@ns_course_schedule.route("/options")
class CourseScheduleOptions(Resource):

    @ns_course_schedule.doc("get_course_schedule_options", description="获取课程表选项")
    @ns_course_schedule.response(200, "成功")
    @requires_permission("view_classes")
    @cached_api(ttl=60)
    def get(self):
        """获取课程表相关选项（班级、科目、节次）"""
        return APIResponse.success(data=get_schedule_options_view())

@ns_course_schedule.route("/check-conflict")
class CourseScheduleConflictCheck(Resource):

    @ns_course_schedule.doc("check_course_schedule_conflict", description="检查课程时间冲突")
    @ns_course_schedule.response(200, "成功")
    @requires_permission("schedule.view")
    def get(self):
        """检查课程时间冲突。支持按班级、教师、教室维度检测。"""
        args = ns_course_schedule.parser.parse_args()
        return APIResponse.success(data=check_schedule_conflict_view(args))

@ns_course_schedule.route("/export")
class CourseScheduleExport(Resource):

    @ns_course_schedule.doc("export_course_schedules", description="导出课程表数据")
    @requires_permission("schedule.view")
    def get(self):
        """导出课程表数据（支持JSON和Excel格式）"""
        args = ns_course_schedule.parser.parse_args()
        class_info_id = args.get("class_info_id")
        export_format = request.args.get("format", "json").lower()
        export_data, filename_prefix = build_schedule_export_data(class_info_id)

        if export_format == "excel":
            wb = Workbook()
            ws = wb.active
            ws.title = "课程表数据"

            # 设置表头样式
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

            # 写入表头
            headers = [
                "班级名称",
                "班级年级",
                "科目名称",
                "星期",
                "节次",
                "节次名称",
                "教师",
                "教室",
                "备注",
                "是否启用",
                "创建时间",
            ]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            # 写入数据
            for row_idx, item in enumerate(export_data, 2):
                ws.cell(row=row_idx, column=1, value=item["class_name"]).border = thin_border
                ws.cell(row=row_idx, column=2, value=item["class_grade"]).border = thin_border
                ws.cell(row=row_idx, column=3, value=item["subject_name"]).border = thin_border
                ws.cell(row=row_idx, column=4, value=item["day_of_week_text"]).border = thin_border
                ws.cell(row=row_idx, column=5, value=item["period_number"]).border = thin_border
                ws.cell(row=row_idx, column=6, value=item["period_name"]).border = thin_border
                ws.cell(row=row_idx, column=7, value=item["teacher_name"]).border = thin_border
                ws.cell(row=row_idx, column=8, value=item["classroom"]).border = thin_border
                ws.cell(row=row_idx, column=9, value=item["description"]).border = thin_border
                ws.cell(row=row_idx, column=10, value=item["is_active"]).border = thin_border
                ws.cell(row=row_idx, column=11, value=item["created_at"]).border = thin_border

            # 调整列宽
            column_widths = [15, 10, 12, 10, 8, 15, 12, 12, 30, 10, 20]
            for i, width in enumerate(column_widths, 1):
                ws.column_dimensions[chr(64 + i) if i <= 26 else f"A{chr(64 + i - 26)}"].width = (
                    width
                )

            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)

            return send_file(
                buf,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True,
                download_name=f"{filename_prefix}.xlsx",
            )
        output = {
            "export_time": datetime.now().isoformat(),
            "total": len(export_data),
            "class_info_id": class_info_id,
            "data": export_data,
        }

        json_str = json.dumps(output, ensure_ascii=False, indent=2)
        buf = io.BytesIO(json_str.encode("utf-8"))
        buf.seek(0)

        return send_file(
            buf,
            mimetype="application/json",
            as_attachment=True,
            download_name=f"{filename_prefix}.json",
        )
