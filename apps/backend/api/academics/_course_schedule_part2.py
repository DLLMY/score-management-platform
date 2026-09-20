# -*- coding: utf-8 -*-
# part of api/academics/course_schedule_routes.py (D2 split)

import json
import io
from flask_restx import Namespace, Resource, fields
from flask import request, send_file
from models import CourseSchedule, ClassInfo, Subject, ClassPeriod, Admin, ImportConfig, get_by_id
from services.academics_service import academics_service
from services.course_schedule_service import (
    get_schedule_list_view,
    get_schedule_by_class_view,
    get_schedule_now_view,
    get_schedule_options_view,
    check_schedule_conflict_view,
    build_schedule_export_data,
    _schedule_dict,
    format_day_of_week,
    get_period_info,
    check_conflicts,
    check_teacher_conflicts,
    check_classroom_conflicts,
)
from utils.permission import requires_permission, get_allowed_classes, get_current_admin
from utils.response import APIResponse
from utils.api_cache_middleware import cached_api, invalidate_cache
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import logging

from api.academics.course_schedule_routes import logger, ns_course_schedule, course_schedule_model, course_schedule_response, _DAY_TEXT_MAP, _validate_text_field, _validate_day_of_week, _validate_period_number, _validate_teacher_role, _check_schedule_update_forbidden, _resolve_schedule_update_fields, _resolve_teacher_name, _schedule_time_changed, _collect_time_conflicts, _check_change_conflicts, _resolve_final_color, _schedule_update_response

@ns_course_schedule.route("/import")
class CourseScheduleImport(Resource):

    @ns_course_schedule.doc("import_course_schedules", description="导入课程表数据")
    @requires_permission("schedule.manage")
    def post(self):
        """从JSON或Excel文件导入课程表数据（支持配置映射）"""
        content_type = request.content_type or ""
        config_id = request.args.get("config_id", type=int)
        strategy_param = request.args.get("conflict_strategy", type=str)

        field_mappings, conflict_strategy, default_values = _load_course_import_config(
            config_id, strategy_param
        )

        import_list = _parse_course_import_input(content_type, field_mappings, default_values)

        success_count = 0
        failed_count = 0
        messages = []
        errors = []
        creates = []
        updates = []

        day_text_map = _DAY_TEXT_MAP

        max_period = ClassPeriod.query.filter_by(is_active=True).count()

        for row_idx, item in enumerate(import_list, start=2):
            try:
                (row_errors, class_info, subject, day_of_week, period_number,
                 teacher_name, classroom) = _validate_course_row(
                    item, day_text_map, max_period
                )
                if row_errors:
                    error_msg = "; ".join(
                        [f'{err["field"]}: {err["message"]}' for err in row_errors]
                    )
                    class_name = item.get("class_name")
                    subject_name = item.get("subject_name")
                    messages.append(
                        {
                            "class_name": class_name or "未知",
                            "subject_name": subject_name or "未知",
                            "action": "failed",
                            "message": error_msg,
                            "row_data": item,
                            "error_fields": [err["field"] for err in row_errors],
                        }
                    )
                    errors.append(
                        {
                            "row": row_idx,
                            "message": error_msg,
                            "row_data": item,
                            "error_fields": [err["field"] for err in row_errors],
                        }
                    )
                    failed_count += 1
                    continue

                conflicts = _detect_course_conflicts(
                    class_info, subject, day_of_week, period_number,
                    teacher_name, classroom
                )
                if conflicts:
                    error_msg = "; ".join([c["message"] for c in conflicts])
                    messages.append(
                        {
                            "class_name": class_info.name,
                            "subject_name": subject.name,
                            "action": "failed",
                            "message": error_msg,
                            "row_data": item,
                            "error_fields": ["conflict"],
                        }
                    )
                    errors.append(
                        {
                            "row": row_idx,
                            "message": error_msg,
                            "row_data": item,
                            "error_fields": ["conflict"],
                        }
                    )
                    failed_count += 1
                    continue

                existing = CourseSchedule.query.filter(
                    CourseSchedule.class_info_id == class_info.id,
                    CourseSchedule.day_of_week == day_of_week,
                    CourseSchedule.period_number == period_number,
                ).first()

                action, payload, message = _resolve_existing_course(
                    existing, item, class_info, subject, day_of_week,
                    period_number, teacher_name, classroom, conflict_strategy
                )
                if action == "skip":
                    messages.append(message)
                    continue
                if action == "error":
                    messages.append(message)
                    errors.append(
                        {
                            "row": row_idx,
                            "message": message["message"],
                            "row_data": item,
                            "error_fields": ["conflict"],
                        }
                    )
                    failed_count += 1
                    continue
                if action == "update":
                    updates.append(payload)
                elif action == "create":
                    creates.append(payload)
                messages.append(message)
                success_count += 1
            except Exception as e:
                error_msg = str(e)
                messages.append(
                    {
                        "class_name": item.get("class_name", "未知"),
                        "subject_name": item.get("subject_name", "未知"),
                        "action": "failed",
                        "message": error_msg,
                        "row_data": item,
                        "error_fields": ["system"],
                    }
                )
                errors.append(
                    {
                        "row": row_idx,
                        "message": error_msg,
                        "row_data": item,
                        "error_fields": ["system"],
                    }
                )
                failed_count += 1

        academics_service.apply_course_schedule_import(creates, updates)
        invalidate_cache("api:/api/course-schedules/*")

        return APIResponse.success(
            data={
                "success": True,
                "total": len(import_list),
                "success_count": success_count,
                "failed_count": failed_count,
                "messages": messages,
            }
        )

def _load_course_import_config(config_id, strategy_param):
    """加载导入配置：指定 config_id 优先，否则取默认配置；返回映射/策略/默认值。"""
    config = None
    if config_id:
        config = get_by_id(ImportConfig, config_id)
    else:
        config = ImportConfig.query.filter(
            ImportConfig.module_name == "course_schedule",
            ImportConfig.is_default,
            ImportConfig.is_active,
        ).first()

    default_mappings = [
        {
            "source_field": "班级名称",
            "target_field": "class_name",
            "field_type": "string",
            "required": True,
            "relation": "class_info",
        },
        {
            "source_field": "科目名称",
            "target_field": "subject_name",
            "field_type": "string",
            "required": True,
            "relation": "subject",
        },
        {
            "source_field": "星期",
            "target_field": "day_of_week",
            "field_type": "string",
            "required": True,
        },
        {
            "source_field": "节次",
            "target_field": "period_number",
            "field_type": "integer",
            "required": True,
        },
        {"source_field": "教师", "target_field": "teacher_name", "field_type": "string"},
        {"source_field": "教室", "target_field": "classroom", "field_type": "string"},
        {"source_field": "备注", "target_field": "description", "field_type": "string"},
        {"source_field": "是否启用", "target_field": "is_active", "field_type": "boolean"},
    ]

    field_mappings = config.field_mappings if config else default_mappings
    conflict_strategy = (
        strategy_param
        if strategy_param and strategy_param in ["skip", "update", "error"]
        else (config.conflict_strategy if config else "update")
    )
    default_values = config.default_values if config else {}
    return field_mappings, conflict_strategy, default_values

def _parse_course_import_multipart(field_mappings, default_values):
    """解析 multipart/form-data 上传：校验文件存在与格式，委托 Excel 解析。"""
    if "file" not in request.files:
        return APIResponse.bad_request(message="请上传文件")

    file = request.files["file"]
    if not file.filename:
        return APIResponse.bad_request(message="请选择文件")

    filename = file.filename.lower()
    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        return _parse_excel_file(file, field_mappings, default_values)
    return APIResponse.bad_request(message="仅支持 .xlsx 或 .xls 格式")

def _parse_excel_file(file, field_mappings, default_values):
    """用 openpyxl 读取工作表，按字段映射构建导入项列表。"""
    from openpyxl import load_workbook

    wb = load_workbook(file)
    ws = wb.active

    headers = []
    for cell in ws[1]:
        headers.append(cell.value)

    col_map = {}
    for idx, header in enumerate(headers):
        if header:
            col_map[header] = idx

    import_list = []
    for row_idx in range(2, ws.max_row + 1):
        row_data = {}
        for header, col_idx in col_map.items():
            row_data[header] = ws.cell(row=row_idx, column=col_idx + 1).value

        mapped_item = _build_mapped_item(row_data, field_mappings, default_values)
        if mapped_item.get("class_name") and mapped_item.get("subject_name"):
            import_list.append(mapped_item)
    return import_list

def _build_mapped_item(row_data, field_mappings, default_values):
    """按字段映射把一行源数据转换为目标字段字典。"""
    mapped_item = {}
    for mapping in field_mappings:
        source_val = row_data.get(mapping["source_field"])
        target_field = mapping["target_field"]
        field_type = mapping.get("field_type", "string")

        if source_val is None:
            if mapping.get("required"):
                break
            source_val = mapping.get("default_value", default_values.get(target_field))

        mapped_item[target_field] = _convert_field_value(source_val, target_field, field_type)
    return mapped_item

def _convert_field_value(source_val, target_field, field_type):
    """按字段类型/目标字段名将源值转换为目标类型。"""
    if field_type == "boolean":
        if isinstance(source_val, str):
            return source_val in ["是", "true", "True", "1"]
        return bool(source_val)
    if field_type == "integer":
        return int(source_val) if source_val else None
    if target_field == "day_of_week" and isinstance(source_val, str):
        return _DAY_TEXT_MAP.get(source_val, 0)
    return source_val

def _parse_course_import_input(content_type, field_mappings, default_values):
    """按 Content-Type 解析导入源（Excel / JSON）为 import_list。"""
    if "multipart/form-data" in content_type:
        result = _parse_course_import_multipart(field_mappings, default_values)
        if isinstance(result, APIResponse):
            return result
        import_list = result
    elif "application/json" in content_type:
        data = request.json
        if not data or "data" not in data:
            return APIResponse.bad_request(message="导入数据格式错误")
        import_list = data["data"]
    else:
        return APIResponse.bad_request(message="不支持的文件格式")
    return import_list

def _validate_course_import_item(item, day_text_map, max_period):
    """单行导入项校验：返回错误列表（空 = 通过）。"""
    row_errors = []

    class_name = item.get("class_name")
    subject_name = item.get("subject_name")
    day_of_week = item.get("day_of_week")
    period_number = item.get("period_number")
    teacher_name = item.get("teacher_name")
    classroom = item.get("classroom")

    err = _validate_text_field(
        class_name, "class_name", 100,
        "班级名称不能为空", "班级名称格式无效，必须为非空字符串",
        "班级名称长度超过限制（最大100字符）", required=True,
    )
    if err:
        row_errors.append(err)

    err = _validate_text_field(
        subject_name, "subject_name", 50,
        "科目名称不能为空", "科目名称格式无效，必须为非空字符串",
        "科目名称长度超过限制（最大50字符）", required=True,
    )
    if err:
        row_errors.append(err)

    err = _validate_day_of_week(day_of_week, day_text_map)
    if err:
        row_errors.append(err)

    err = _validate_period_number(period_number, max_period)
    if err:
        row_errors.append(err)

    if teacher_name:
        err = _validate_text_field(
            teacher_name, "teacher_name", 50,
            None, "教师姓名格式无效，必须为非空字符串",
            "教师姓名长度超过限制（最大50字符）", required=False,
        )
        if err:
            row_errors.append(err)
        else:
            err = _validate_teacher_role(teacher_name)
            if err:
                row_errors.append(err)

    if classroom:
        err = _validate_text_field(
            classroom, "classroom", 50,
            None, "教室名称格式无效，必须为非空字符串",
            "教室名称长度超过限制（最大50字符）", required=False,
        )
        if err:
            row_errors.append(err)

    return row_errors

def _validate_course_row(item, day_text_map, max_period):
    """Validate one course-schedule import row and resolve its entities.
    Mirrors the original post() validation block exactly.
    """
    row_errors = _validate_course_import_item(item, day_text_map, max_period)
    class_name = item.get("class_name")
    subject_name = item.get("subject_name")
    class_info = None
    subject = None
    if class_name:
        class_info = ClassInfo.query.filter_by(name=class_name.strip()).first()
        if not class_info:
            row_errors.append(
                {
                    "field": "class_name",
                    "message": f'班级 "{class_name}" 在系统中不存在',
                }
            )
    if subject_name:
        subject = Subject.query.filter_by(name=subject_name.strip()).first()
        if not subject:
            row_errors.append(
                {
                    "field": "subject_name",
                    "message": f'科目 "{subject_name}" 在系统中不存在',
                }
            )
    day_of_week = item["day_of_week"]
    if isinstance(day_of_week, str):
        day_of_week = day_text_map.get(day_of_week, 0)
    period_number = item["period_number"]
    teacher_name = item.get("teacher_name")
    classroom = item.get("classroom")
    return row_errors, class_info, subject, day_of_week, period_number, teacher_name, classroom

def _detect_course_conflicts(class_info, subject, day_of_week, period_number,
                             teacher_name, classroom):
    """Collect schedule/teacher/classroom conflicts for one resolved row."""
    conflicts = []
    conflicts.extend(check_conflicts(class_info.id, day_of_week, period_number))
    if teacher_name:
        conflicts.extend(
            check_teacher_conflicts(teacher_name, day_of_week, period_number)
        )
    if classroom:
        conflicts.extend(
            check_classroom_conflicts(classroom, day_of_week, period_number)
        )
    return conflicts

def _resolve_existing_course(existing, item, class_info, subject, day_of_week,
                             period_number, teacher_name, classroom, conflict_strategy):
    """Decide the outcome for one resolved row given the existing schedule.
    Returns (action, payload, message) where action is one of
    skip/update/error/create.
    """
    if existing:
        if conflict_strategy == "skip":
            return "skip", None, {
                "class_name": class_info.name,
                "subject_name": subject.name,
                "action": "skipped",
                "message": (
                    f"{class_info.name} {format_day_of_week(day_of_week)}"
                    f"第{period_number}节课程已存在，已跳过"
                ),
            }
        if conflict_strategy == "update":
            return "update", (existing.id, {
                "subject_id": subject.id,
                "teacher_name": teacher_name or existing.teacher_name,
                "classroom": classroom or existing.classroom,
                "description": item.get("description", existing.description),
                "color": item.get("color", existing.color),
                "is_active": item.get("is_active", existing.is_active),
            }), {
                "class_name": class_info.name,
                "subject_name": subject.name,
                "action": "updated",
                "message": f"{class_info.name} {format_day_of_week(day_of_week)}第{period_number}节课程已更新",
            }
        elif conflict_strategy == "error":
            error_msg = f"{class_info.name} {format_day_of_week(day_of_week)}第{period_number}节课程已存在，与导入数据冲突"
            return "error", None, {
                "class_name": class_info.name,
                "subject_name": subject.name,
                "action": "failed",
                "message": error_msg,
                "row_data": item,
                "error_fields": ["conflict"],
            }
    return "create", {
        "class_info_id": class_info.id,
        "subject_id": subject.id,
        "day_of_week": day_of_week,
        "period_number": period_number,
        "teacher_name": teacher_name,
        "classroom": classroom,
        "description": item.get("description"),
        "color": item.get("color", subject.color),
        "is_active": item.get("is_active", True),
    }, {
        "class_name": class_info.name,
        "subject_name": subject.name,
        "action": "created",
        "message": f"{class_info.name} {format_day_of_week(day_of_week)}第{period_number}节课程已创建",
    }
