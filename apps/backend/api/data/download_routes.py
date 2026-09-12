import logging

import json
import io
import openpyxl
from flask import Blueprint, request, send_file
from utils.decorators import safe_handle

logger = logging.getLogger(__name__)

"""
文件下载路由模块
"""
download_bp = Blueprint("download", __name__)


@download_bp.route("/api/scores/template/download")
def download_score_template():
    """
    下载成绩导入Excel模板
    参数:
    - exam_id: 考试ID（可选）
    - class_name: 班级名称（可选）
    """
    from models import Exam, User, get_by_id
    from utils.permission import requires_permission

    @requires_permission("system.settings")
    @safe_handle(message="生成模板失败", default_status=500)
    def generate_template():
        exam_id = request.args.get("exam_id", type=int)
        class_name = request.args.get("class_name")
        class_id = request.args.get("class_id", type=int)
        exam = None
        if exam_id:
            exam = get_by_id(Exam, exam_id)
        subjects = _resolve_template_subjects(exam)
        students = _query_template_students(User, class_id, class_name)
        output = _build_score_template_workbook(subjects, students)
        filename = f"score_import_template_{class_name or 'all'}.xlsx"
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )

    return generate_template()


def _resolve_template_subjects(exam):
    """解析模板科目列表；无考试或科目为空时回退为语数英。"""
    if exam and exam.subjects:
        return json.loads(exam.subjects) if isinstance(exam.subjects, str) else exam.subjects
    return ["语文", "数学", "英语"]


def _query_template_students(user_model, class_id, class_name):
    """按 class_id > class_name > 全部 的优先级取学生列表。"""
    if class_id:
        return (
            user_model.query.filter(user_model.class_info_id == class_id)
            .order_by(user_model.card_id)
            .all()
        )
    if class_name:
        return user_model.query.filter_by(class_name=class_name).order_by(user_model.card_id).all()
    return user_model.query.order_by(user_model.class_name, user_model.card_id).all()


def _build_score_template_workbook(subjects, students):
    """构建成绩导入模板工作簿（成绩导入表 + 填写说明表），返回内存缓冲。"""
    wb = openpyxl.Workbook()
    score_sheet = wb.active
    score_sheet.title = "成绩导入"
    _fill_score_sheet(score_sheet, subjects, students)
    notes_sheet = wb.create_sheet(title="填写说明")
    _fill_notes_sheet(notes_sheet)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _fill_score_sheet(sheet, subjects, students):
    """写入表头、学生基础信息（科目分数列留空）并设置列宽。"""
    headers = ["学号", "姓名", "班级"] + subjects
    for col_idx, header in enumerate(headers, 1):
        sheet.cell(row=1, column=col_idx, value=header)
    for row_idx, student in enumerate(students, 2):
        sheet.cell(row=row_idx, column=1, value=student.card_id)
        sheet.cell(row=row_idx, column=2, value=student.name)
        sheet.cell(row=row_idx, column=3, value=student.class_name)
        # 科目分数列留空，由教师填写
        for col_offset in range(len(subjects)):
            sheet.cell(row=row_idx, column=4 + col_offset, value="")
    sheet.column_dimensions["A"].width = 15
    sheet.column_dimensions["B"].width = 10
    sheet.column_dimensions["C"].width = 15
    for i in range(len(subjects)):
        col_letter = openpyxl.utils.get_column_letter(4 + i)
        sheet.column_dimensions[col_letter].width = 12


_NOTES_ROWS = [
    ["列名", "说明", "填写方式", "示例"],
    ["学号", "学生的学号，系统自动填入，请勿修改", "系统自动", "202401001"],
    ["姓名", "学生姓名，系统自动填入，仅作参考", "系统自动", "张三"],
    ["班级", "班级名称，系统自动填入", "系统自动", "高一(1)班"],
    ["科目", "科目名称，对应考试科目", "系统自动", "语文"],
    ["分数", "学生成绩，教师必须填写，必须为数字(0-100)", "教师填写", "85"],
]


def _fill_notes_sheet(sheet):
    """写入填写说明表内容并设置列宽。"""
    for row_idx, row_data in enumerate(_NOTES_ROWS, 1):
        for col_idx, cell_value in enumerate(row_data, 1):
            sheet.cell(row=row_idx, column=col_idx, value=cell_value)
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 40
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 15
